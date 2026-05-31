from PIL import Image
import openpyxl
from openpyxl.styles import PatternFill

new_wb= openpyxl.Workbook()
ws=new_wb.active
img=Image.open(r"OIP.jpg")
img.resize((400,400))

for i in range(40):
    for j in range (40):
        r_total = g_total = b_total = 0
        count = 0

        for x in range(10*i, 10*(i+1)):
            for y in range(10*j, 10*(j+1)):
                r, g, b = img.getpixel((y, x))  # دقت کن: x = row, y = column
                r_total += r
                g_total += g
                b_total += b
                count += 1

        r_avg = r_total // count
        g_avg = g_total // count
        b_avg = b_total // count

        # تبدیل به hex برای رنگ سلول
        hex_color = f"{r_avg:02X}{g_avg:02X}{b_avg:02X}"

        # اعمال رنگ پس‌زمینه
        cell = ws.cell(row=i+1, column=j+1)
        cell.fill = PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")

new_wb.save(r"image.xlsx")