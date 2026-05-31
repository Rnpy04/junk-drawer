import os
from pdf2image import convert_from_path
from pptx import Presentation
from pptx.util import Inches

pdf_path = "your_file.pdf"
pptx_path = "output_presentation.pptx"

images = convert_from_path(pdf_path, dpi=200)

prs = Presentation()

for img in images:
    img_path = "temp_slide.png"
    img.save(img_path, "PNG")

    slide_layout = prs.slide_layouts[6] 
    slide = prs.slides.add_slide(slide_layout)

    slide.shapes.add_picture(img_path, Inches(0), Inches(0),
                             width=prs.slide_width, height=prs.slide_height)

    os.remove(img_path)

prs.save(pptx_path)
print("saved")
